"""
Comprehensive Flask Web Application
Integrates Google Scholar Research Agent and Task Assigner

Features:
1. Google Scholar research analysis with marketing focus
2. Task assignment from Google Docs to Trello
3. Real-time progress tracking
4. Document generation
5. User-friendly web interface
"""

import os
import json
import logging
import threading
import time
from datetime import datetime
from typing import Dict, Any
import uuid
import re # Added for AI Content Generator

# Configure logging first
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

from flask import Flask, render_template, request, jsonify, redirect, url_for, flash, session, send_file
from docx import Document # Added for AI Content Generator
from docx.enum.text import WD_ALIGN_PARAGRAPH # Added for AI Content Generator

# Import research agent (optional - will handle import errors)
try:
    from google_scholar_research_agent import GoogleScholarResearchAgent
    RESEARCH_AGENT_AVAILABLE = True
    logger.info("✅ Google Scholar Research Agent imported successfully")
except ImportError as e:
    logger.warning(f"⚠️ Google Scholar Research Agent not available: {e}")
    GoogleScholarResearchAgent = None
    RESEARCH_AGENT_AVAILABLE = False

# Import task assignment module (optional - will handle import errors)
try:
    import sys
    sys.path.append('task assigner')
    from task_assigner import TaskAssigner
    TASK_ASSIGNMENT_AVAILABLE = True
    logger.info("✅ Task Assigner imported successfully")
except ImportError as e:
    logger.warning(f"⚠️ Task Assignment module not available: {e}")
    TaskAssigner = None
    TASK_ASSIGNMENT_AVAILABLE = False

# Flask app configuration
app = Flask(__name__)
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY', 'your-secret-key-here')
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024  # 16MB max file size

# Global storage for task status and results
task_status = {}
task_results = {}

# Global instances
research_agent = None
task_assigner = None

def initialize_services():
    """Initialize both research agent and task assigner"""
    global research_agent, task_assigner

    # Initialize Research Agent
    if RESEARCH_AGENT_AVAILABLE:
        try:
            # Initialize research agent with service account authentication (no OAuth popups)
            research_agent = GoogleScholarResearchAgent(
                openai_api_key=os.getenv("OPENAI_API_KEY"),
                google_credentials_path="credentials.json"  # Will try service account first
            )
            
            # Configure Bright Data
            research_agent.configure_bright_data_api(
                enable_bright_data=True,
                api_key=os.getenv("BRIGHT_DATA_API_KEY", "your_api_key_here")
            )
            
            # Configure PDF processing
            research_agent.configure_paper_download(
                enable_download=True,
                max_download_size_mb=50,
                download_timeout=30,
                max_parallel_downloads=3
            )
            
            logger.info("✅ Research agent initialized successfully (service account auth - no OAuth popups)")
        except Exception as e:
            logger.error(f"❌ Failed to initialize research agent: {e}")
            research_agent = None
    
    # Initialize Task Assigner
    if TASK_ASSIGNMENT_AVAILABLE:
        try:
            task_assigner = TaskAssigner()
            logger.info("✅ Task assigner initialized successfully")
        except Exception as e:
            logger.error(f"❌ Failed to initialize task assigner: {e}")
            task_assigner = None

def get_configuration_status():
    """Get comprehensive configuration status for both research and task assignment"""
    
    # Check for service account specifically
    service_account_available = (
        os.path.exists('service-account.json') or 
        os.path.exists('task assigner/service-account.json') or
        bool(os.getenv('GOOGLE_SERVICE_ACCOUNT_PATH'))
    )
    
    oauth_token_available = os.path.exists('token.json')
    
    # Research Agent Configuration (Service Account - no OAuth popups)
    research_config = {
        'openai_available': bool(os.getenv('OPENAI_API_KEY')),
        'bright_data_available': bool(os.getenv('BRIGHT_DATA_API_KEY')),
        'google_credentials_available': service_account_available or oauth_token_available,
        'service_account_available': service_account_available,
        'oauth_token_available': oauth_token_available
    }
    
    # Task Assignment Configuration  
    task_config = {
        'google_docs_available': (
            os.path.exists('credentials.json') or 
            os.path.exists('task assigner/service-account.json') or
            os.path.exists('token.json')
        ),
        'trello_available': bool(os.getenv('TRELLO_API_KEY') and os.getenv('TRELLO_TOKEN')),
        'team_available': os.path.exists('task assigner/team_members.json')
    }
    
    return research_config, task_config

@app.route('/')
def index():
    """Main dashboard page"""
    research_config, task_config = get_configuration_status()
    
    services_status = {
        'research_agent': research_agent is not None,
        'task_assigner': task_assigner is not None,
        'research_agent_available': RESEARCH_AGENT_AVAILABLE,
        'task_assignment_available': TASK_ASSIGNMENT_AVAILABLE
    }
    
    return render_template('index.html', 
                         services=services_status,
                         research_config=research_config,
                         task_config=task_config)

# ================================
# RESEARCH AGENT ROUTES
# ================================

@app.route('/research')
def research_form():
    """Research analysis form page"""
    if not RESEARCH_AGENT_AVAILABLE:
        flash('Research agent is not available. Please check your setup.', 'error')
        return redirect(url_for('index'))

    research_config, _ = get_configuration_status()
    return render_template('research_form.html', config_status=research_config)

@app.route('/research/analyze', methods=['POST'])
def research_analyze():
    """Handle research analysis request - now OAuth-free"""
    if not research_agent:
        return jsonify({'error': 'Research agent not available'}), 500
        
    try:
        # Get form data
        google_file_id = request.form.get('google_file_id', '').strip()
        max_papers = int(request.form.get('max_papers', 10))
        analysis_focus = request.form.get('analysis_focus', 'general')
        output_formats = request.form.getlist('output_formats')
        
        # Validate inputs
        if not google_file_id:
            return jsonify({'error': 'Google File ID is required'}), 400
            
        # Generate unique task ID
        task_id = str(uuid.uuid4())
        
        # Initialize task status
        task_status[task_id] = {
            'id': task_id,
            'status': 'starting',
            'progress': 0,
            'message': 'Initializing research analysis...',
            'created_at': datetime.now().isoformat(),
            'steps': [
                'Extracting keywords from Google Doc',
                'Searching Google Scholar',
                'Filtering and ranking papers',
                'Downloading PDFs',
                'Analyzing content',
                'Marketing relevance analysis',
                'Generating document'
            ],
            'current_step': 0,
            'total_steps': 7,
            'details': [],
            'type': 'research',
            'papers_found': 0,
            'papers_processed': 0,
            'pdfs_downloaded': 0,
            'bright_data_stats': {}
        }
        
        # Configure research agent based on form options
        research_agent.max_papers_per_query = max_papers
        
        # Start background analysis from Google file
        thread = threading.Thread(
            target=run_research_from_google_file,
            args=(task_id, google_file_id, output_formats, analysis_focus)
        )
        
        thread.daemon = True
        thread.start()
        
        return jsonify({
            'task_id': task_id,
            'status': 'started',
            'message': 'Research analysis started (service account authentication)',
            'estimated_time': '5-15 minutes'
        })
        
    except Exception as e:
        logger.error(f"Error starting research analysis: {e}")
        return jsonify({'error': f'Failed to start analysis: {str(e)}'}), 500

def run_research_from_google_file(task_id: str, google_file_id: str, output_formats: list, analysis_focus: str):
    """Run complete research pipeline from Google file"""
    try:
        # Update status
        update_task_status(task_id, 'running', 'Starting research pipeline from Google file...', 10)
        
        # Step 1: Extract keywords from Google file using service account
        update_task_status(task_id, 'running', 'Extracting keywords from Google file (service account auth)...', 20, 1)
        keywords_data = research_agent.extract_keywords_from_google_file(google_file_id)
        
        if "error" in keywords_data:
            error_msg = keywords_data["error"]
            logger.error(f"Google file extraction failed: {error_msg}")
            
            # Provide helpful error message to user
            if "authentication" in error_msg.lower() or "service-account" in error_msg.lower():
                user_error = "Google authentication failed. Please ensure service-account.json file is properly configured."
            elif "shared" in error_msg.lower() or "permission" in error_msg.lower():
                user_error = f"Cannot access Google file. Please share the document with: sebastiancastano@phonic-goods-317118.iam.gserviceaccount.com"
            else:
                user_error = f"Failed to extract keywords from Google file: {error_msg}"
            
            update_task_status(task_id, 'error', user_error, 20)
            return
        
        # Step 2: Create sophisticated search queries from enhanced keywords
        search_queries = []
        
        # Use search_queries if available (these are optimized for academic paper searches)
        if keywords_data.get("search_queries"):
            search_queries.extend(keywords_data.get("search_queries", [])[:4])
        
        # Add primary keywords and research topics as backup
        search_queries.extend(keywords_data.get("primary_keywords", [])[:3])
        search_queries.extend(keywords_data.get("research_topics", [])[:2])
        
        # Add academic concepts for more sophisticated searches
        search_queries.extend(keywords_data.get("academic_concepts", [])[:2])
        
        # Remove duplicates while preserving order
        seen = set()
        unique_queries = []
        for query in search_queries:
            if query.lower() not in seen:
                unique_queries.append(query)
                seen.add(query.lower())
        
        search_queries = unique_queries[:8]  # Limit to 8 total queries for efficiency
        
        logger.info(f"[SEARCH] Using sophisticated search queries: {search_queries}")
        update_task_status(task_id, 'running', f'Searching Google Scholar with {len(search_queries)} sophisticated queries...', 30, 2)
        papers = research_agent.search_google_scholar(search_queries)
        
        task_status[task_id]['papers_found'] = len(papers)
        
        if not papers:
            update_task_status(task_id, 'error', 'No papers found from Google Scholar search', 30)
            return
        
        # Continue with common pipeline
        run_common_research_pipeline(task_id, papers, keywords_data, search_queries, output_formats, analysis_focus)
        
    except Exception as e:
        logger.error(f"Error in research pipeline from Google file: {e}")
        update_task_status(task_id, 'error', f'Pipeline error: {str(e)}', task_status[task_id]['progress'])


def run_research_oauth_free(task_id: str, google_file_id: str = None, manual_keywords: str = None, output_formats: list = None, analysis_focus: str = 'general'):
    """Run research pipeline without Google OAuth - supports both file ID and manual keywords"""
    try:
        output_formats = output_formats or ['word']
        
        # Step 1: Get keywords (OAuth-free approach)
        if manual_keywords:
            update_task_status(task_id, 'running', 'Processing manual keywords...', 15, 1)
            # Process manual keywords
            keywords_list = [kw.strip() for kw in manual_keywords.split(',') if kw.strip()]
            keywords_data = {
                "primary_keywords": keywords_list[:5],
                "secondary_keywords": keywords_list[5:10] if len(keywords_list) > 5 else [],
                "research_topics": keywords_list[:3],
                "methodologies": ["analysis", "research", "study"]
            }
        elif google_file_id:
            update_task_status(task_id, 'running', 'Note: Google file access requires OAuth (disabled). Using fallback approach...', 15, 1)
            # Fallback: Use generic keywords based on common research terms
            keywords_data = {
                "primary_keywords": ["research", "analysis", "study", "methodology", "findings"],
                "secondary_keywords": ["investigation", "examination", "evaluation", "assessment", "review"],
                "research_topics": ["academic research", "scientific study"],
                "methodologies": ["analysis", "research methodology", "systematic review"]
            }
            # Add a note about the limitation
            update_task_status(task_id, 'running', 'Using generic research keywords since Google OAuth is disabled', 20, 1)
        else:
            update_task_status(task_id, 'error', 'No keywords provided', 15)
            return

        # Step 2: Create sophisticated search queries from enhanced keywords
        search_queries = []
        
        # Use search_queries if available (these are optimized for academic paper searches)
        if keywords_data.get("search_queries"):
            search_queries.extend(keywords_data.get("search_queries", [])[:4])
        
        # Add primary keywords and research topics as backup
        search_queries.extend(keywords_data.get("primary_keywords", [])[:3])
        search_queries.extend(keywords_data.get("research_topics", [])[:2])
        
        # Add academic concepts for more sophisticated searches
        search_queries.extend(keywords_data.get("academic_concepts", [])[:2])
        
        # Remove duplicates while preserving order
        seen = set()
        unique_queries = []
        for query in search_queries:
            if query.lower() not in seen:
                unique_queries.append(query)
                seen.add(query.lower())
        
        search_queries = unique_queries[:8]  # Limit to 8 total queries for efficiency
        
        logger.info(f"[SEARCH] Using sophisticated search queries: {search_queries}")
        update_task_status(task_id, 'running', f'Searching Google Scholar with {len(search_queries)} sophisticated queries...', 30, 2)
        papers = research_agent.search_google_scholar(search_queries)
        
        task_status[task_id]['papers_found'] = len(papers)
        
        if not papers:
            update_task_status(task_id, 'error', 'No papers found from Google Scholar search', 30)
            return
        
        # Continue with common pipeline
        run_common_research_pipeline(task_id, papers, keywords_data, search_queries, output_formats, analysis_focus)
        
    except Exception as e:
        logger.error(f"Error in OAuth-free research pipeline: {e}")
        update_task_status(task_id, 'error', f'Pipeline error: {str(e)}', task_status[task_id]['progress'])


def run_common_research_pipeline(task_id: str, papers: list, keywords_data: dict, search_queries: list, output_formats: list, analysis_focus: str):
    """Common research pipeline for both input methods"""
    try:
        # Step 3: Filter and rank papers
        update_task_status(task_id, 'running', 'Filtering and ranking papers by relevance...', 40, 3)
        filtered_papers = research_agent.filter_and_rank_papers(papers, " ".join(search_queries))
        
        task_status[task_id]['papers_processed'] = len(filtered_papers)
        
        # Step 4: Download and process PDFs with timeout and progress tracking
        update_task_status(task_id, 'running', 'Downloading and processing PDFs...', 50, 4)
        
        processed_papers = []
        pdfs_downloaded = 0
        
        try:
            logger.info(f"[TASK-{task_id}] Starting PDF download for {len(filtered_papers)} papers")
            
            # Use threading-based timeout that works on all platforms
            import threading
            import queue
            
            def process_papers_with_timeout():
                """Process papers with individual timeouts"""
                try:
                    papers_result = []
                    pdf_count = 0
                    
                    # Process papers with progress updates
                    for i, paper in enumerate(filtered_papers):
                        try:
                            # Calculate progress within step 4 (50-55%)
                            step_progress = 50 + int((i / len(filtered_papers)) * 5)
                            update_task_status(task_id, 'running', f'Processing paper {i+1}/{len(filtered_papers)}: {paper.title[:50]}...', step_progress, 4)
                            logger.info(f"[TASK-{task_id}] Processing paper {i+1}/{len(filtered_papers)}: {paper.title[:50]}")
                            
                            # Try to download PDF with timeout per paper
                            def download_single_paper():
                                try:
                                    if hasattr(research_agent, '_download_and_process_single_pdf'):
                                        return research_agent._download_and_process_single_pdf(paper)
                                    else:
                                        # Fallback: try to enhance the paper's abstract if no PDF method available
                                        if hasattr(research_agent, '_enhance_abstract_with_ai'):
                                            research_agent._enhance_abstract_with_ai(paper)
                                        return paper
                                except Exception as e:
                                    logger.warning(f"[TASK-{task_id}] Individual paper processing failed: {e}")
                                    return paper
                            
                            # Use threading to implement timeout per paper
                            result_queue = queue.Queue()
                            
                            def worker():
                                try:
                                    result = download_single_paper()
                                    result_queue.put(('success', result))
                                except Exception as e:
                                    result_queue.put(('error', str(e)))
                            
                            worker_thread = threading.Thread(target=worker)
                            worker_thread.daemon = True
                            worker_thread.start()
                            worker_thread.join(timeout=45)  # 45 seconds per paper
                            
                            if worker_thread.is_alive():
                                logger.warning(f"[TASK-{task_id}] Timeout processing paper: {paper.title[:50]}")
                                papers_result.append(paper)  # Use original paper
                            else:
                                try:
                                    status, result = result_queue.get_nowait()
                                    if status == 'success' and result:
                                        papers_result.append(result)
                                        if hasattr(result, 'full_text') and result.full_text:
                                            pdf_count += 1
                                            logger.info(f"[TASK-{task_id}] Successfully processed: {paper.title[:50]}")
                                    else:
                                        papers_result.append(paper)
                                        if status == 'error':
                                            logger.warning(f"[TASK-{task_id}] Error processing {paper.title[:50]}: {result}")
                                except queue.Empty:
                                    papers_result.append(paper)
                                    
                        except Exception as e:
                            logger.error(f"[TASK-{task_id}] Error processing paper {paper.title[:50]}: {e}")
                            papers_result.append(paper)  # Keep original paper
                            continue
                            
                    return papers_result, pdf_count
                    
                except Exception as e:
                    logger.error(f"[TASK-{task_id}] Critical error in paper processing: {e}")
                    return filtered_papers, 0
            
            # Run the processing with overall timeout (5 minutes)
            processing_queue = queue.Queue()
            
            def processing_worker():
                try:
                    result = process_papers_with_timeout()
                    processing_queue.put(('success', result))
                except Exception as e:
                    processing_queue.put(('error', str(e)))
            
            processing_thread = threading.Thread(target=processing_worker)
            processing_thread.daemon = True
            processing_thread.start()
            processing_thread.join(timeout=300)  # 5 minutes total timeout
            
            if processing_thread.is_alive():
                logger.error(f"[TASK-{task_id}] PDF download process timed out after 5 minutes")
                processed_papers = filtered_papers  # Fallback to original papers
                pdfs_downloaded = 0
            else:
                try:
                    status, result = processing_queue.get_nowait()
                    if status == 'success':
                        processed_papers, pdfs_downloaded = result
                    else:
                        logger.error(f"[TASK-{task_id}] Processing failed: {result}")
                        processed_papers = filtered_papers
                        pdfs_downloaded = 0
                except queue.Empty:
                    logger.warning(f"[TASK-{task_id}] No result received from processing")
                    processed_papers = filtered_papers
                    pdfs_downloaded = 0
                    
        except Exception as e:
            logger.error(f"[TASK-{task_id}] Critical error in PDF processing setup: {e}")
            processed_papers = filtered_papers
            pdfs_downloaded = 0
        
        # Ensure we have papers to continue with
        if not processed_papers:
            logger.warning(f"[TASK-{task_id}] No papers processed, using filtered papers")
            processed_papers = filtered_papers
            
        task_status[task_id]['pdfs_downloaded'] = pdfs_downloaded
        logger.info(f"[TASK-{task_id}] PDF processing complete: {pdfs_downloaded} PDFs downloaded from {len(processed_papers)} papers")
        
        # Step 5: Analyze research content
        update_task_status(task_id, 'running', 'Analyzing research content with AI...', 60, 5)
        analysis_data = research_agent.analyze_research_content(processed_papers)
                
        # Step 6: Analyze marketing relevance
        update_task_status(task_id, 'running', 'Analyzing marketing relevance and business applications...', 70, 6)
        marketing_analysis = research_agent.analyze_marketing_relevance(keywords_data, processed_papers)
        
        # Step 7: Generate documents based on selected output formats
        update_task_status(task_id, 'running', 'Generating comprehensive research documents...', 80, 7)
        
        context = f"Research Focus: {analysis_focus}, Keywords: {', '.join(search_queries)}"
        
        # Initialize document results
        documents_created = {}
        
        # Create Word document if requested
        if 'word' in output_formats:
            update_task_status(task_id, 'running', 'Creating Word document...', 85, 7)
            doc_filename = research_agent.create_research_document(
                analysis_data, 
                context,
                marketing_analysis
            )
            if doc_filename:
                documents_created['word'] = {
                    'filename': doc_filename,
                    'type': 'download',
                    'path': doc_filename
                }
                logger.info(f"[DOC] Word document created: {doc_filename}")
        
        # Create Google Doc if requested
        if 'google_doc' in output_formats:
            update_task_status(task_id, 'running', 'Creating Google Document...', 90, 7)
            google_doc_result = research_agent.create_google_document(
                analysis_data,
                context,
                marketing_analysis
            )
            if 'error' not in google_doc_result:
                documents_created['google_doc'] = {
                    'document_id': google_doc_result['document_id'],
                    'document_url': google_doc_result['document_url'],
                    'view_url': google_doc_result['view_url'],
                    'edit_url': google_doc_result['edit_url'],
                    'title': google_doc_result['title'],
                    'type': 'link'
                }
                logger.info(f"[DOC] Google Document created: {google_doc_result['document_url']}")
            else:
                logger.error(f"[ERROR] Google Document creation failed: {google_doc_result['error']}")
                # Store the error for display
                documents_created['google_doc_error'] = google_doc_result['error']
        
        # Create PDF if requested (convert from Word)
        if 'pdf' in output_formats and 'word' in documents_created:
            update_task_status(task_id, 'running', 'Converting to PDF format...', 95, 7)
            # For now, just note that PDF creation would go here
            # In a full implementation, you'd convert the Word doc to PDF
            documents_created['pdf_note'] = 'PDF conversion available upon request'
        
        # Update Bright Data statistics
        task_status[task_id]['bright_data_stats'] = research_agent.get_bright_data_stats()
        
        # Store results
        task_results[task_id] = {
            'documents_created': documents_created,
            'papers_found': len(papers),
            'papers_analyzed': len(filtered_papers),
            'papers_processed': len(processed_papers),
            'pdfs_downloaded': pdfs_downloaded,
            'keywords_data': keywords_data,
            'marketing_analysis': marketing_analysis,
            'academic_analysis': analysis_data,
            'bright_data_stats': research_agent.get_bright_data_stats(),
            'analysis_focus': analysis_focus,
            'search_queries': search_queries,
            'output_formats': output_formats
        }
        
        # Final completion
        update_task_status(task_id, 'completed', 'Research analysis completed successfully!', 100, 8)
        
    except Exception as e:
        logger.error(f"Error in common research pipeline: {e}")
        update_task_status(task_id, 'error', f'Analysis error: {str(e)}', task_status[task_id]['progress'])

def update_task_status(task_id: str, status: str, message: str, progress: int, step: int = None):
    """Update task status with detailed information"""
    if task_id in task_status:
        task_status[task_id].update({
            'status': status,
            'message': message,
            'progress': progress,
            'last_update': time.time()
        })
        if step is not None:
            task_status[task_id]['current_step'] = step
        
        # Add to details log
        task_status[task_id]['details'].append({
            'timestamp': datetime.now().strftime('%H:%M:%S'),
            'message': message,
            'progress': progress
        })
        
        logger.info(f"Task {task_id}: {message} ({progress}%)")

@app.route('/research/progress/<task_id>')
def research_progress(task_id):
    """Research analysis progress page"""
    if task_id not in task_status:
        flash('Task not found', 'error')
        return redirect(url_for('research_form'))
    
    return render_template('progress.html', task_id=task_id, task_info=task_status[task_id])

@app.route('/research/results/<task_id>')
def research_results(task_id):
    """Research analysis results page"""
    if task_id not in task_status:
        flash('Task not found', 'error')
        return redirect(url_for('research_form'))
    
    if task_status[task_id]['status'] != 'completed':
        return redirect(url_for('research_progress', task_id=task_id))
    
    results = task_results.get(task_id, {})
    return render_template('results.html', 
                         task_id=task_id, 
                         task_info=task_status[task_id],
                         results=results)

@app.route('/api/research/status/<task_id>')
def api_research_status(task_id):
    """API endpoint to get research analysis status"""
    if task_id not in task_status:
        return jsonify({'error': 'Task not found'}), 404
    
    status_data = task_status[task_id].copy()
    
    # Add results if completed
    if status_data['status'] == 'completed' and task_id in task_results:
        status_data['results'] = task_results[task_id]
    
    return jsonify(status_data)

@app.route('/download/<filename>')
def download_file(filename):
    """Download generated research document"""
    try:
        file_path = os.path.join(os.getcwd(), filename)
        if os.path.exists(file_path):
            return send_file(file_path, as_attachment=True, download_name=filename)
        else:
            flash('File not found', 'error')
            return redirect(url_for('index'))
    except Exception as e:
        logger.error(f"Error downloading file: {e}")
        flash('Download failed', 'error')
        return redirect(url_for('index'))

# ================================
# TASK ASSIGNMENT ROUTES
# ================================

@app.route('/tasks')
def task_assignment_form():
    """Task assignment form page"""
    if not TASK_ASSIGNMENT_AVAILABLE:
        flash('Task assignment is not available. Please check your setup.', 'error')
        return redirect(url_for('index'))
    
    # Check configuration
    config_status = check_task_assignment_config()
    return render_template('task_assignment_form.html', config=config_status)

@app.route('/tasks/assign', methods=['POST'])
def assign_tasks():
    """Start task assignment process"""
    if not TASK_ASSIGNMENT_AVAILABLE or not task_assigner:
        return jsonify({'error': 'Task assignment not available'}), 500
    
    try:
        google_file_id = request.form.get('google_file_id')
        debug_mode = request.form.get('debug_mode') == 'true'
        
        if not google_file_id:
            return jsonify({'error': 'Google File ID is required'}), 400
        
        # Generate unique task ID
        task_id = str(uuid.uuid4())
        
        # Log debug mode if enabled
        if debug_mode:
            logger.info(f"[TASK-{task_id}] 🐛 Debug mode enabled for task assignment")
        
        # Initialize task status
        task_status[task_id] = {
            'id': task_id,
            'status': 'starting',
            'message': 'Initializing task assignment...',
            'progress': 0,
            'start_time': time.time(),
            'last_update': time.time(),
            'created_at': datetime.now().isoformat(),
            'steps': [
                'Connecting to Google Docs',
                'Extracting document content',
                'Analyzing content with AI',
                'Identifying and parsing tasks',
                'Assigning tasks to team members',
                'Creating Trello cards'
            ],
            'current_step': 0,
            'total_steps': 6,
            'details': [],
            'type': 'task_assignment',
            'debug_mode': debug_mode,
            'tasks_found': 0,
            'tasks_assigned': 0,
            'cards_created': 0
        }
        
        # Start background task assignment
        thread = threading.Thread(
            target=run_task_assignment_process,
            args=(task_id, google_file_id, debug_mode)
        )
        thread.daemon = True
        thread.start()
        
        return jsonify({
            'task_id': task_id,
            'status': 'started',
            'message': 'Task assignment started',
            'estimated_time': '3-8 minutes'
        })
        
    except Exception as e:
        logger.error(f"Error starting task assignment: {e}")
        return jsonify({'error': f'Failed to start task assignment: {str(e)}'}), 500

def run_task_assignment_process(task_id: str, google_file_id: str, debug_mode: bool = False):
    """Run task assignment process in background using TaskAssigner"""
    try:
        debug_prefix = "🐛 [DEBUG]" if debug_mode else ""
        logger.info(f"[TASK-{task_id}] {debug_prefix} Starting task assignment process for Google Doc: {google_file_id}")
        
        if not task_assigner:
            raise Exception("Task assigner not initialized")
            
        if debug_mode:
            logger.info(f"[TASK-{task_id}] 🐛 Debug mode active - Enhanced logging enabled")
            logger.info(f"[TASK-{task_id}] 🐛 Task assigner instance: {type(task_assigner)}")
            logger.info(f"[TASK-{task_id}] 🐛 Google File ID: {google_file_id}")
        
        # Step 1: Connecting to Google Docs
        update_task_status(task_id, 'running', 'Connecting to Google Docs...', 15, 0)
        logger.info(f"[TASK-{task_id}] Step 1: Connecting to Google Docs")
        
        # Step 2: Extracting document content
        update_task_status(task_id, 'running', 'Extracting document content...', 30, 1)
        logger.info(f"[TASK-{task_id}] Step 2: Extracting document content")
        
        # Extract text from Google Doc
        try:
            content = task_assigner.extract_text_from_google_doc(google_file_id)
            logger.info(f"[TASK-{task_id}] Extracted {len(content)} characters from document")
            
            if not content or len(content.strip()) < 50:
                raise Exception("Document appears to be empty or too short to process")
                
        except Exception as e:
            logger.error(f"[TASK-{task_id}] Error extracting Google Doc content: {e}")
            raise Exception(f"Failed to extract document content: {str(e)}")
        
        # Step 3: Analyzing content with AI
        update_task_status(task_id, 'running', 'Analyzing content with AI...', 50, 2)
        logger.info(f"[TASK-{task_id}] Step 3: Analyzing content with OpenAI")
        
        # Process content with OpenAI to extract tasks
        try:
            tasks = task_assigner.process_content_with_openai(content)
            task_status[task_id]['tasks_found'] = len(tasks)
            logger.info(f"[TASK-{task_id}] Extracted {len(tasks)} tasks from content")
            
            if not tasks:
                raise Exception("No actionable tasks found in the document")
                
        except Exception as e:
            logger.error(f"[TASK-{task_id}] Error processing content with OpenAI: {e}")
            raise Exception(f"Failed to analyze content: {str(e)}")
        
        # Step 4: Identifying and parsing tasks
        update_task_status(task_id, 'running', 'Identifying and parsing tasks...', 65, 3)
        logger.info(f"[TASK-{task_id}] Step 4: Processing {len(tasks)} identified tasks")
        
        # Step 5: Assigning tasks to team members
        update_task_status(task_id, 'running', 'Assigning tasks to team members...', 80, 4)
        logger.info(f"[TASK-{task_id}] Step 5: Assigning tasks to team members")
        
        # Assign tasks and track assignments
        assignments = {}
        trello_cards = []
        errors = []
        
        for i, task in enumerate(tasks):
            try:
                # Assign to team member
                assignee = task_assigner.assign_task_to_team_member(task)
                task['assignee'] = assignee
                
                # Update assignment counts
                if assignee not in assignments:
                    assignments[assignee] = 0
                assignments[assignee] += 1
                
                logger.info(f"[TASK-{task_id}] Assigned task '{task.get('title', 'Untitled')}' to {assignee}")
                
            except Exception as e:
                error_msg = f"Error assigning task '{task.get('title', 'Unknown')}': {str(e)}"
                logger.error(f"[TASK-{task_id}] {error_msg}")
                errors.append(error_msg)
        
        task_status[task_id]['tasks_assigned'] = len([t for t in tasks if t.get('assignee')])
        
        # Step 6: Creating Trello cards
        update_task_status(task_id, 'running', 'Creating Trello cards...', 90, 5)
        logger.info(f"[TASK-{task_id}] Step 6: Creating Trello cards")
        
        # Create Trello cards
        for task in tasks:
            if not task.get('assignee'):
                continue
                
            try:
                card_info = task_assigner.create_trello_card(task, task['assignee'])
                if card_info:
                    trello_cards.append({
                        'task': task['title'],
                        'assignee': task['assignee'],
                        'card_id': card_info['id'],
                        'card_url': card_info['url']
                    })
                    logger.info(f"[TASK-{task_id}] Created Trello card for '{task['title']}'")
                else:
                    error_msg = f"Failed to create Trello card for '{task.get('title', 'Unknown')}'"
                    logger.warning(f"[TASK-{task_id}] {error_msg}")
                    errors.append(error_msg)
                    
            except Exception as e:
                error_msg = f"Error creating Trello card for '{task.get('title', 'Unknown')}': {str(e)}"
                logger.error(f"[TASK-{task_id}] {error_msg}")
                errors.append(error_msg)
        
        task_status[task_id]['cards_created'] = len(trello_cards)
        
        # Prepare results
        results = {
            'total_tasks': len(tasks),
            'tasks': tasks,
            'assignments': assignments,
            'trello_cards': trello_cards,
            'errors': errors,
            'trello_board_url': f"https://trello.com/b/{task_assigner.trello_board_id}" if hasattr(task_assigner, 'trello_board_id') else None
        }
        
        # Store results
        task_results[task_id] = results
        
        # Complete the task
        update_task_status(task_id, 'completed', 'Task assignment completed successfully!', 100, 6)
        logger.info(f"[TASK-{task_id}] Task assignment completed successfully. Created {len(trello_cards)} Trello cards.")
        
        # Log summary
        logger.info(f"[TASK-{task_id}] Summary - Tasks: {len(tasks)}, Assigned: {len([t for t in tasks if t.get('assignee')])}, Cards: {len(trello_cards)}, Errors: {len(errors)}")
        
    except Exception as e:
        logger.error(f"[TASK-{task_id}] Critical error in task assignment: {e}")
        error_msg = f'Task assignment error: {str(e)}'
        current_progress = task_status.get(task_id, {}).get('progress', 0)
        update_task_status(task_id, 'error', error_msg, current_progress)

def check_task_assignment_config():
    """Check task assignment configuration"""
    return {
        'google_docs_ready': os.path.exists('credentials.json'),
        'trello_ready': bool(os.getenv('TRELLO_API_KEY')),
        'team_config_ready': os.path.exists('task assigner/team_members.json'),
        'ready': os.path.exists('credentials.json') and bool(os.getenv('TRELLO_API_KEY'))
    }

@app.route('/api/tasks/config')
def api_task_config():
    """API endpoint to check task assignment configuration"""
    if not TASK_ASSIGNMENT_AVAILABLE:
        return jsonify({
            'error': 'Task assignment module not available',
            'ready': False,
            'missing': ['task_assignment_module']
        }), 500
    
    return jsonify(check_task_assignment_config())

@app.route('/api/tasks/status/<task_id>')
def api_task_status(task_id):
    """API endpoint to get task assignment status"""
    if task_id not in task_status:
        return jsonify({'error': 'Task not found'}), 404
    
    status_data = task_status[task_id].copy()
    
    # Add results if completed
    if status_data['status'] == 'completed' and task_id in task_results:
        status_data['results'] = task_results[task_id]
    
    return jsonify(status_data)

@app.route('/api/task/status/<task_id>')
def api_task_status_alt(task_id):
    """Alternative API endpoint to match frontend expectation"""
    return api_task_status(task_id)

@app.route('/tasks/results/<task_id>')
def task_assignment_results(task_id):
    """Task assignment results page"""
    if task_id not in task_status:
        flash('Task not found', 'error')
        return redirect(url_for('task_assignment_form'))
    
    if task_status[task_id]['status'] != 'completed':
        return redirect(url_for('task_assignment_progress', task_id=task_id))
    
    results = task_results.get(task_id, {})
    task_info = task_status[task_id]
    
    # Calculate some statistics
    stats = {
        'total_tasks': results.get('total_tasks', 0),
        'team_assignments': results.get('assignments', {}),
        'trello_cards_created': len(results.get('trello_cards', [])),
        'errors_count': len(results.get('errors', [])),
        'duration': int(time.time() - task_info.get('start_time', time.time()))
    }
    
    return render_template('task_results.html', 
                         task_id=task_id, 
                         task=task_info,
                         results=results,
                         stats=stats)

@app.route('/tasks/progress/<task_id>')
def task_assignment_progress(task_id):
    """Task assignment progress page"""
    if task_id not in task_status:
        flash('Task not found', 'error')
        return redirect(url_for('task_assignment_form'))
    
    return render_template('task_progress.html', task_id=task_id, task=task_status[task_id])

# Routes for AI Content Generator
@app.route('/ai-content')
def ai_content_generator_form():
    """Render AI content generator form"""
    try:
        return render_template('ai_content_generator.html')
    except Exception as e:
        logger.error(f"Error rendering AI content generator form: {e}")
        flash(f'Error loading AI content generator: {str(e)}', 'error')
        return redirect(url_for('index'))

@app.route('/ai-content/generate', methods=['POST'])
def start_ai_content_generation():
    """Start AI content generation process"""
    try:
        # Get form data
        google_file_id = request.form.get('google_file_id', '').strip()
        content_type = request.form.get('content_type', 'blog_post')
        content_tone = request.form.get('content_tone', 'professional')
        content_length = request.form.get('content_length', 'medium')
        target_audience = request.form.get('target_audience', 'general')
        additional_instructions = request.form.get('additional_instructions', '').strip()
        
        # Validate input
        if not google_file_id:
            return jsonify({'error': 'Google File ID is required'}), 400
        
        # Extract actual file ID if full URL is provided
        if 'docs.google.com' in google_file_id:
            match = re.search(r'/document/d/([a-zA-Z0-9-_]+)', google_file_id)
            if match:
                google_file_id = match.group(1)
            else:
                return jsonify({'error': 'Invalid Google Docs URL format'}), 400
        
        # Generate unique task ID
        task_id = str(uuid.uuid4())
        
        # Initialize task status
        task_status[task_id] = {
            'status': 'starting',
            'message': 'Initializing AI content generation...',
            'progress': 0,
            'step': 0,
            'total_steps': 6,
            'start_time': time.time(),
            'content_type': content_type,
            'content_tone': content_tone,
            'content_length': content_length,
            'target_audience': target_audience
        }
        
        # Start background task
        thread = threading.Thread(
            target=run_ai_content_generation_process,
            args=(task_id, google_file_id, content_type, content_tone, content_length, target_audience, additional_instructions)
        )
        thread.daemon = True
        thread.start()
        
        logger.info(f"[AI-CONTENT-{task_id}] Started AI content generation for Google Doc: {google_file_id}")
        
        return jsonify({
            'task_id': task_id,
            'status': 'started',
            'message': 'AI content generation started',
            'estimated_time': '2-5 minutes'
        })
        
    except Exception as e:
        logger.error(f"Error starting AI content generation: {e}")
        return jsonify({'error': f'Failed to start AI content generation: {str(e)}'}), 500

def run_ai_content_generation_process(task_id: str, google_file_id: str, content_type: str, 
                                    content_tone: str, content_length: str, target_audience: str, 
                                    additional_instructions: str):
    """Run AI content generation process in background"""
    try:
        logger.info(f"[AI-CONTENT-{task_id}] Starting AI content generation process for Google Doc: {google_file_id}")
        
        if not research_agent:
            raise Exception("Research agent not initialized")
        
        # Step 1: Extract keywords from Google file
        update_task_status(task_id, 'running', 'Extracting keywords from Google file...', 15, 1)
        logger.info(f"[AI-CONTENT-{task_id}] Step 1: Extracting keywords from Google file")
        
        keywords_data = research_agent.extract_keywords_from_google_file(google_file_id)
        
        if "error" in keywords_data:
            raise Exception(f"Failed to extract keywords: {keywords_data['error']}")
        
        # Step 2: Perform research based on keywords
        update_task_status(task_id, 'running', 'Conducting research based on extracted keywords...', 35, 2)
        logger.info(f"[AI-CONTENT-{task_id}] Step 2: Conducting research")
        
        # Create search queries from keywords
        search_queries = []
        search_queries.extend(keywords_data.get("primary_keywords", [])[:3])
        search_queries.extend(keywords_data.get("research_topics", [])[:2])
        
        # Search Google Scholar
        papers = research_agent.search_google_scholar(search_queries)
        
        if not papers:
            logger.warning(f"[AI-CONTENT-{task_id}] No papers found, proceeding with keyword-based content")
            papers = []
        
        # Filter and rank papers
        filtered_papers = research_agent.filter_and_rank_papers(papers, " ".join(search_queries))
        
        # Step 3: Process papers and extract insights
        update_task_status(task_id, 'running', 'Processing research papers and extracting insights...', 55, 3)
        logger.info(f"[AI-CONTENT-{task_id}] Step 3: Processing papers")
        
        processed_papers = research_agent.download_and_process_pdfs(filtered_papers)
        
        # Step 4: Analyze research for content generation
        update_task_status(task_id, 'running', 'Analyzing research for content generation...', 70, 4)
        logger.info(f"[AI-CONTENT-{task_id}] Step 4: Analyzing research")
        
        analysis_data = research_agent.analyze_research_content(processed_papers)
        marketing_analysis = research_agent.analyze_marketing_relevance(keywords_data, processed_papers)
        
        # Step 5: Generate AI content
        update_task_status(task_id, 'running', 'Generating AI content based on research...', 85, 5)
        logger.info(f"[AI-CONTENT-{task_id}] Step 5: Generating AI content")
        
        # Generate the AI content
        content_result = generate_ai_content_from_research(
            keywords_data, analysis_data, marketing_analysis, processed_papers,
            content_type, content_tone, content_length, target_audience, additional_instructions
        )
        
        # Step 6: Create final output document
        update_task_status(task_id, 'running', 'Creating final content document...', 95, 6)
        logger.info(f"[AI-CONTENT-{task_id}] Step 6: Creating output document")
        
        # Create document with the generated content
        final_document = create_ai_content_document(
            content_result, keywords_data, analysis_data, content_type, content_tone, content_length
        )
        
        # Store results
        task_results[task_id] = {
            'content_generated': content_result,
            'final_document': final_document,
            'research_summary': {
                'papers_found': len(papers),
                'papers_analyzed': len(filtered_papers),
                'papers_processed': len(processed_papers),
                'keywords_extracted': len(keywords_data.get("primary_keywords", [])),
            },
            'generation_params': {
                'content_type': content_type,
                'content_tone': content_tone,
                'content_length': content_length,
                'target_audience': target_audience,
                'additional_instructions': additional_instructions
            },
            'bright_data_stats': research_agent.get_bright_data_stats() if research_agent else {}
        }
        
        # Final completion
        update_task_status(task_id, 'completed', 'AI content generation completed successfully!', 100, 6)
        logger.info(f"[AI-CONTENT-{task_id}] AI content generation completed successfully")
        
    except Exception as e:
        logger.error(f"[AI-CONTENT-{task_id}] Error in AI content generation: {e}")
        update_task_status(task_id, 'error', f'Error: {str(e)}', 0, 0)

def generate_ai_content_from_research(keywords_data, analysis_data, marketing_analysis, papers, 
                                    content_type, content_tone, content_length, target_audience, additional_instructions):
    """Generate AI content based on research data"""
    try:
        if not research_agent or not research_agent.client:
            raise Exception("OpenAI client not available")
        
        # Prepare research context
        research_context = {
            'keywords': keywords_data,
            'academic_analysis': analysis_data,
            'marketing_analysis': marketing_analysis,
            'papers_summary': [
                {
                    'title': paper.title,
                    'authors': paper.authors,
                    'abstract': paper.abstract,
                    'key_findings': getattr(paper, 'key_findings', []),
                    'business_relevance': getattr(paper, 'business_relevance', ''),
                    'citations': paper.citations
                } for paper in papers[:5]  # Top 5 papers
            ]
        }
        
        # Content generation prompts based on type
        prompts = {
            'blog_post': create_blog_post_prompt,
            'article': create_article_prompt,
            'social_media': create_social_media_prompt,
            'marketing_copy': create_marketing_copy_prompt,
            'whitepaper': create_whitepaper_prompt,
            'email_campaign': create_email_campaign_prompt,
            'press_release': create_press_release_prompt,
            'case_study': create_case_study_prompt
        }
        
        if content_type not in prompts:
            content_type = 'blog_post'  # Default fallback
        
        # Generate the content prompt
        content_prompt = prompts[content_type](
            research_context, content_tone, content_length, target_audience, additional_instructions
        )
        
        # Generate content with OpenAI
        response = research_agent.client.chat.completions.create(
            model=research_agent.model_name,
            messages=[
                {
                    "role": "system", 
                    "content": f"You are an expert content creator specializing in {content_type.replace('_', ' ')} creation. Create engaging, well-researched content that effectively communicates complex information to the target audience."
                },
                {
                    "role": "user", 
                    "content": content_prompt
                }
            ],
            temperature=0.7,
            max_tokens=4000
        )
        
        generated_content = response.choices[0].message.content
        
        return {
            'content': generated_content,
            'content_type': content_type,
            'word_count': len(generated_content.split()),
            'research_papers_used': len(papers),
            'generation_successful': True
        }
        
    except Exception as e:
        logger.error(f"Error generating AI content: {e}")
        return {
            'content': f"Error generating content: {str(e)}",
            'content_type': content_type,
            'word_count': 0,
            'research_papers_used': 0,
            'generation_successful': False,
            'error': str(e)
        }

# Content generation prompt functions
def create_blog_post_prompt(research_context, tone, length, audience, additional_instructions):
    """Create blog post generation prompt"""
    word_counts = {'short': '800-1200', 'medium': '1200-1800', 'long': '1800-2500'}
    target_words = word_counts.get(length, '1200-1800')
    
    return f"""
Create a comprehensive blog post based on the following research data:

RESEARCH KEYWORDS: {research_context['keywords'].get('primary_keywords', [])}
RESEARCH FINDINGS: {research_context['academic_analysis'].get('marketing_summary', 'Research analysis available')}
MARKETING INSIGHTS: {research_context['marketing_analysis'].get('executive_summary', 'Marketing analysis available')}

TOP RESEARCH PAPERS:
{chr(10).join([f"- {paper['title']} by {', '.join(paper['authors'][:2])}" for paper in research_context['papers_summary'][:3]])}

CONTENT REQUIREMENTS:
- Type: Blog Post
- Tone: {tone}
- Length: {target_words} words
- Target Audience: {audience}
- Additional Instructions: {additional_instructions}

Create an engaging, well-structured blog post that:
1. Has a compelling headline
2. Includes an engaging introduction
3. Incorporates research findings naturally
4. Provides actionable insights
5. Includes subheadings for readability
6. Concludes with key takeaways
7. Cites research appropriately

Make the content valuable for the target audience while maintaining the specified tone and length.
"""

def create_article_prompt(research_context, tone, length, audience, additional_instructions):
    """Create article generation prompt"""
    word_counts = {'short': '1000-1500', 'medium': '1500-2500', 'long': '2500-4000'}
    target_words = word_counts.get(length, '1500-2500')
    
    return f"""
Create a comprehensive article based on the following research data:

RESEARCH KEYWORDS: {research_context['keywords'].get('primary_keywords', [])}
ACADEMIC ANALYSIS: {research_context['academic_analysis'].get('marketing_summary', 'Research analysis available')}
MARKETING INSIGHTS: {research_context['marketing_analysis'].get('executive_summary', 'Marketing analysis available')}

RESEARCH PAPERS ANALYZED:
{chr(10).join([f"- {paper['title']}" for paper in research_context['papers_summary'][:5]])}

CONTENT REQUIREMENTS:
- Type: In-depth Article
- Tone: {tone}
- Length: {target_words} words
- Target Audience: {audience}
- Additional Instructions: {additional_instructions}

Create a well-researched, authoritative article that:
1. Has a compelling title
2. Includes an executive summary
3. Provides detailed analysis of research findings
4. Offers multiple perspectives
5. Includes data and statistics
6. Provides actionable recommendations
7. Has proper citations and references

Ensure the article demonstrates thought leadership and provides deep insights for the target audience.
"""

def create_social_media_prompt(research_context, tone, length, audience, additional_instructions):
    """Create social media content generation prompt"""
    platforms = {'short': 'Twitter threads', 'medium': 'LinkedIn posts', 'long': 'Multi-platform campaign'}
    platform_type = platforms.get(length, 'LinkedIn posts')
    
    return f"""
Create engaging social media content based on the following research data:

RESEARCH KEYWORDS: {research_context['keywords'].get('primary_keywords', [])}
KEY FINDINGS: {research_context['academic_analysis'].get('marketing_summary', 'Research insights available')}
MARKETING INSIGHTS: {research_context['marketing_analysis'].get('executive_summary', 'Marketing analysis available')}

CONTENT REQUIREMENTS:
- Type: {platform_type}
- Tone: {tone}
- Target Audience: {audience}
- Additional Instructions: {additional_instructions}

Create compelling social media content that:
1. Grabs attention immediately
2. Incorporates research insights
3. Uses engaging visuals concepts
4. Includes relevant hashtags
5. Encourages engagement
6. Provides value to followers
7. Maintains brand voice

Provide multiple post variations and engagement strategies.
"""

def create_marketing_copy_prompt(research_context, tone, length, audience, additional_instructions):
    """Create marketing copy generation prompt"""
    copy_types = {'short': 'Email subject lines and ads', 'medium': 'Landing page copy', 'long': 'Complete campaign copy'}
    copy_type = copy_types.get(length, 'Landing page copy')
    
    return f"""
Create persuasive marketing copy based on the following research data:

RESEARCH KEYWORDS: {research_context['keywords'].get('primary_keywords', [])}
MARKET INSIGHTS: {research_context['marketing_analysis'].get('executive_summary', 'Marketing analysis available')}
RESEARCH FINDINGS: {research_context['academic_analysis'].get('marketing_summary', 'Research insights available')}

CONTENT REQUIREMENTS:
- Type: {copy_type}
- Tone: {tone}
- Target Audience: {audience}
- Additional Instructions: {additional_instructions}

Create high-converting marketing copy that:
1. Addresses target audience pain points
2. Incorporates research-backed benefits
3. Uses persuasive language
4. Includes strong calls-to-action
5. Builds credibility with research citations
6. Creates urgency and desire
7. Optimizes for conversion

Provide multiple variations and A/B testing suggestions.
"""

def create_whitepaper_prompt(research_context, tone, length, audience, additional_instructions):
    """Create whitepaper generation prompt"""
    return f"""
Create a comprehensive whitepaper based on the following research data:

RESEARCH FOUNDATION: {research_context['keywords'].get('primary_keywords', [])}
ACADEMIC RESEARCH: {research_context['academic_analysis'].get('marketing_summary', 'Comprehensive research analysis available')}
MARKET ANALYSIS: {research_context['marketing_analysis'].get('executive_summary', 'Market insights available')}

RESEARCH PAPERS:
{chr(10).join([f"- {paper['title']} (Citations: {paper['citations']})" for paper in research_context['papers_summary'][:5]])}

CONTENT REQUIREMENTS:
- Type: Professional Whitepaper
- Tone: {tone}
- Target Audience: {audience}
- Additional Instructions: {additional_instructions}

Create an authoritative whitepaper that:
1. Has an executive summary
2. Defines the problem/opportunity
3. Presents research methodology
4. Analyzes findings comprehensively
5. Provides strategic recommendations
6. Includes charts and data concepts
7. Has proper citations and bibliography

Structure as a professional, downloadable resource that establishes thought leadership.
"""

def create_email_campaign_prompt(research_context, tone, length, audience, additional_instructions):
    """Create email campaign generation prompt"""
    return f"""
Create an email campaign series based on the following research data:

RESEARCH INSIGHTS: {research_context['keywords'].get('primary_keywords', [])}
MARKETING ANALYSIS: {research_context['marketing_analysis'].get('executive_summary', 'Marketing insights available')}
RESEARCH FINDINGS: {research_context['academic_analysis'].get('marketing_summary', 'Research analysis available')}

CONTENT REQUIREMENTS:
- Type: Email Campaign Series
- Tone: {tone}
- Target Audience: {audience}
- Additional Instructions: {additional_instructions}

Create a multi-email campaign that:
1. Has compelling subject lines
2. Nurtures leads through value
3. Incorporates research insights
4. Builds trust and authority
5. Includes clear CTAs
6. Segments by audience needs
7. Optimizes for deliverability

Provide 3-5 emails in sequence with specific purposes and goals.
"""

def create_press_release_prompt(research_context, tone, length, audience, additional_instructions):
    """Create press release generation prompt"""
    return f"""
Create a press release based on the following research data:

RESEARCH BREAKTHROUGH: {research_context['keywords'].get('primary_keywords', [])}
RESEARCH FINDINGS: {research_context['academic_analysis'].get('marketing_summary', 'Significant research findings available')}
MARKET IMPLICATIONS: {research_context['marketing_analysis'].get('executive_summary', 'Market impact analysis available')}

CONTENT REQUIREMENTS:
- Type: Press Release
- Tone: {tone}
- Target Audience: {audience} (Media and stakeholders)
- Additional Instructions: {additional_instructions}

Create a newsworthy press release that:
1. Has a compelling headline
2. Follows proper press release format
3. Leads with the most important news
4. Incorporates research credibility
5. Includes relevant quotes
6. Provides contact information
7. Optimizes for media pickup

Make it newsworthy and shareable by journalists and industry publications.
"""

def create_case_study_prompt(research_context, tone, length, audience, additional_instructions):
    """Create case study generation prompt"""
    return f"""
Create a compelling case study based on the following research data:

RESEARCH BASIS: {research_context['keywords'].get('primary_keywords', [])}
ANALYTICAL FRAMEWORK: {research_context['academic_analysis'].get('marketing_summary', 'Research methodology and findings available')}
BUSINESS APPLICATIONS: {research_context['marketing_analysis'].get('executive_summary', 'Business impact analysis available')}

SUPPORTING RESEARCH:
{chr(10).join([f"- {paper['title']}" for paper in research_context['papers_summary'][:3]])}

CONTENT REQUIREMENTS:
- Type: Business Case Study
- Tone: {tone}
- Target Audience: {audience}
- Additional Instructions: {additional_instructions}

Create a detailed case study that:
1. Presents a clear challenge/opportunity
2. Describes the research-based approach
3. Details implementation methodology
4. Presents measurable results
5. Includes lessons learned
6. Provides actionable takeaways
7. Uses data visualization concepts

Structure as a professional case study that demonstrates real-world application of research insights.
"""

def create_ai_content_document(content_result, keywords_data, analysis_data, content_type, content_tone, content_length):
    """Create a document with the generated AI content"""
    try:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"ai_content_{content_type}_{timestamp}.docx"
        
        # Create Word document
        doc = Document()
        
        # Title
        title = doc.add_heading(f'AI-Generated {content_type.replace("_", " ").title()}', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Metadata
        doc.add_heading('Content Details', level=1)
        doc.add_paragraph(f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_paragraph(f"Content Type: {content_type.replace('_', ' ').title()}")
        doc.add_paragraph(f"Tone: {content_tone.title()}")
        doc.add_paragraph(f"Length: {content_length.title()}")
        doc.add_paragraph(f"Word Count: {content_result.get('word_count', 'Unknown')}")
        doc.add_paragraph(f"Research Papers Used: {content_result.get('research_papers_used', 0)}")
        
        # Keywords used
        doc.add_heading('Research Keywords', level=1)
        primary_keywords = keywords_data.get('primary_keywords', [])
        for keyword in primary_keywords[:10]:
            doc.add_paragraph(f"• {keyword}")
        
        # Generated content
        doc.add_heading('Generated Content', level=1)
        
        # Split content into paragraphs for better formatting
        content_text = content_result.get('content', '')
        paragraphs = content_text.split('\n\n')
        
        for paragraph in paragraphs:
            if paragraph.strip():
                doc.add_paragraph(paragraph.strip())
        
        # Research basis
        doc.add_heading('Research Foundation', level=1)
        research_summary = analysis_data.get('marketing_summary', 'Research analysis provided the foundation for this content.')
        doc.add_paragraph(research_summary)
        
        # Save document
        doc.save(filename)
        
        return {
            'filename': filename,
            'type': 'download',
            'path': filename,
            'word_count': content_result.get('word_count', 0),
            'success': True
        }
        
    except Exception as e:
        logger.error(f"Error creating AI content document: {e}")
        return {
            'filename': '',
            'type': 'error',
            'path': '',
            'error': str(e),
            'success': False
        }

@app.route('/ai-content/progress/<task_id>')
def ai_content_progress(task_id):
    """Show AI content generation progress"""
    try:
        if task_id not in task_status:
            return render_template('error.html', 
                                 error_title="Task Not Found",
                                 error_message="The requested AI content generation task was not found.")
        
        return render_template('ai_content_progress.html', task_id=task_id)
        
    except Exception as e:
        logger.error(f"Error showing AI content progress: {e}")
        return render_template('error.html', 
                             error_title="Progress Error",
                             error_message=f"Error displaying progress: {str(e)}")

@app.route('/ai-content/results/<task_id>')
def ai_content_results(task_id):
    """Show AI content generation results"""
    try:
        if task_id not in task_status:
            return render_template('error.html', 
                                 error_title="Task Not Found",
                                 error_message="The requested AI content generation task was not found.")
        
        status = task_status[task_id]
        results = task_results.get(task_id, {})
        
        if status['status'] != 'completed':
            return redirect(url_for('ai_content_progress', task_id=task_id))
        
        return render_template('ai_content_results.html', 
                             task_id=task_id, 
                             status=status, 
                             results=results)
        
    except Exception as e:
        logger.error(f"Error showing AI content results: {e}")
        return render_template('error.html', 
                             error_title="Results Error",
                             error_message=f"Error displaying results: {str(e)}")

@app.route('/api/ai-content/status/<task_id>')
def api_ai_content_status(task_id):
    """API endpoint for AI content generation status"""
    try:
        if task_id not in task_status:
            return jsonify({'error': 'Task not found'}), 404
        
        status = task_status[task_id].copy()
        
        # Add timing information
        if 'start_time' in status:
            elapsed_time = time.time() - status['start_time']
            status['elapsed_time'] = round(elapsed_time, 2)
            
            # Estimate remaining time based on progress
            if status['progress'] > 0:
                estimated_total = elapsed_time / (status['progress'] / 100)
                remaining = max(0, estimated_total - elapsed_time)
                status['estimated_remaining'] = round(remaining, 2)
        
        return jsonify(status)
        
    except Exception as e:
        logger.error(f"Error getting AI content status: {e}")
        return jsonify({'error': str(e)}), 500

# ================================
# ERROR HANDLERS
# ================================

@app.errorhandler(404)
def not_found_error(error):
    return render_template('error.html', 
                         error_code=404, 
                         error_message="Page not found"), 404

@app.errorhandler(500)
def internal_error(error):
    return render_template('error.html', 
                         error_code=500, 
                         error_message="Internal server error"), 500

# ================================
# STARTUP
# ================================

# Initialize services when the app starts
with app.app_context():
    initialize_services()

if __name__ == '__main__':
    import os
    port = int(os.environ.get('PORT', 5000))
    debug = os.environ.get('FLASK_ENV') != 'production'
    app.run(host='0.0.0.0', port=port, debug=debug) 